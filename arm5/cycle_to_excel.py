"""Turn one recorded pick and drop from arm5.html into an Excel workbook.

The page's recorder (window.ARM5_RECORD, or the "Record one pick and drop" button for a CSV) runs one full
vision cycle with the page's own 1/240 s control step. This writes its JSON to sheets:

  Columns      every column with its unit and meaning
  Cycle        every control step: program line, tool and joints, torques, clearances, network inputs/outputs
  Steps        one row per program line: timing, how it moved, peak torque per joint, closest clearances
  GNN calls    only the steps where the pose GNN ran: its 5 inputs and 5 outputs
  PINN         only the steps the motion PINN drove: its 13 inputs and 5 outputs
  Motion GNN calls  only the steps where the motion GNN ran: its 12 inputs and 5 outputs
  (a sheet is left out when its network did not run in the recorded cycle)
  Camera       the image the vision network read (its input), one row per pixel, and what it answered

    python cycle_to_excel.py cycle.json --out pick_and_drop.xlsx
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
HEAD_FILL = PatternFill("solid", fgColor="DDE6F0")
UNIT_FILL = PatternFill("solid", fgColor="F1F4F8")

MEANING = {
    "t": "Time since the cycle started (end of this control step; one step = 1/240 s).",
    "step_no": "Number of the program line being run.", "step": "Name of the program line.",
    "op": "PTP = joint move to a pose, LIN = straight tool line, LOOK = camera, GRIP = gripper.",
    "motion": "What produced the joint angles on this step.",
    "tool_x": "Tool point, horizontal position (the shoulder is at x 0).", "tool_z": "Tool point height above the floor.",
    "tool_angle": "Tool direction (0 = along +x, -90 = straight down).",
    "q_A": "Joint angle.", "qd_A": "Joint speed (finite difference over one step).",
    "qdd_A": "Joint acceleration (finite difference over two steps).",
    "torque_A": "Motor torque needed on this step (Newton-Euler, with the held part).",
    "gripper_opening": "Gripper jaw opening (0 = closed, 1 = fully open).",
    "holding_part": "1 while the part is in the gripper.", "payload": "Mass the arm is carrying on this step.",
    "part_x": "Part position, horizontal.", "part_z": "Part centre height.",
    "self_clearance": "Smallest gap between two parts of the arm that can meet, or between the arm and its pedestal (negative = collision).",
    "self_closest_parts": "Which two parts that smallest gap is between.",
    "part_clearance_from_arm": "While carrying: smallest gap between the part and the arm's links, housings and pedestal.",
    "lowest_link_above_floor": "Lowest edge of links 1-4 above the floor.",
    "gnn_called": "1 when the pose GNN ran on this step (every target, and every 12th step while tracking a line).",
    "gnn_target": "What the GNN was asked for: tool point, tool angle and payload.",
    "gnn_in": "THE GNN'S INPUT (5 numbers, unitless).",
    "gnn_out": "THE GNN'S OUTPUT: joint angles it proposes (before the page refines them).",
    "pinn_active": "1 when the motion PINN set the joint angles on this step.",
    "pinn_s": "Normalized time in the move, t / T.", "pinn_T": "Duration of the move.", "pinn_payload": "Payload the move is planned for.",
    "pinn_in": "THE PINN'S INPUT (13 numbers, unitless).",
    "pinn_out": "THE PINN'S OUTPUT: path correction N; q(s) = q0 + (q1 - q0) b(s) + w(s) N.",
    "mgnn_called": "1 when the motion GNN ran on this step (planning a move, then every 12th step for the next moment ahead).",
    "mgnn_s": "The moment of the move it was asked about, t / T (ahead of the current one).", "mgnn_T": "Duration of the move.",
    "mgnn_in": "THE MOTION GNN'S INPUT (12 numbers, unitless): start pose, target and payload, moment, duration.",
    "mgnn_out": "THE MOTION GNN'S OUTPUT: each joint's end angle e for that moment; q(s) = q0 + b(s) (e - q0).",
    "mgnn_q": "The joint angle that output gives at that moment (with the page's end-pose refinement spread over the move).",
    "vision_called": "1 on the step the camera picture was taken and the vision network ran.",
    "vision_out": "THE VISION NETWORK'S OUTPUT: where it found the part.", "vision_true": "Where the part really was.",
    "vision_error": "Distance between found and true position.",
}


def meaning(name):
    for k in sorted(MEANING, key=len, reverse=True):
        if name == k or name.startswith(k):
            return MEANING[k]
    return ""


def write(ws, head, units, rows, widths=None):
    ws.append(head)
    ws.append(units)
    for r in rows:
        ws.append([round(v, 6) if isinstance(v, float) else (None if v == "" else v) for v in r])
    for c in range(1, len(head) + 1):
        for row, fill, bold in ((1, HEAD_FILL, True), (2, UNIT_FILL, False)):
            cell = ws.cell(row, c)
            cell.fill, cell.font, cell.alignment = fill, Font(bold=bold, italic=not bold), Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(head[c - 1], max(10, min(18, len(head[c - 1]) + 2)))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 32, 45
    ws.freeze_panes = "B3"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("record")
    ap.add_argument("--out", default="pick_and_drop.xlsx")
    args = ap.parse_args()
    d = json.loads(Path(args.record).read_text())
    head, units, rows = d["head"], d["units"], d["rows"]
    col = {h: i for i, h in enumerate(head)}
    wb = Workbook()

    ws = wb.active
    ws.title = "Columns"
    ws.append([f"One pick and drop recorded from arm5.html: motion {d['motion']}, IK {d['solver']}, part {d['kg']} kg, "
               f"{len(rows)} control steps of 1/240 s ({rows[-1][0]:.2f} s)"])
    ws.append([])
    ws.append(["Column", "Unit", "Meaning"])
    for h, u in zip(head, units):
        ws.append([h, u, meaning(h)])
    ws["A1"].font = Font(bold=True, size=12)
    for c in "ABC":
        ws[f"{c}3"].font, ws[f"{c}3"].fill = Font(bold=True), HEAD_FILL
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 26, 34, 110

    write(wb.create_sheet("Cycle"), head, units, rows, {"step": 24, "motion": 26, "self_closest_parts": 34})

    groups = []                                            # consecutive runs of one program line (line 1 comes again at the end)
    for r in rows:
        if not groups or groups[-1][-1][col["step_no"]] != r[col["step_no"]]:
            groups.append([])
        groups[-1].append(r)
    sh = (["step_no", "step", "op", "motion", "start", "end", "duration"] + [f"peak_torque_A{i}" for i in range(1, 6)]
          + [f"peak_torque_A{i}_of_limit" for i in range(1, 6)] + ["min_self_clearance", "closest_parts", "min_part_clearance",
                                                                   "gnn_calls", "pinn_steps", "motion_gnn_calls"])
    su = ["-", "-", "-", "-", "s", "s", "s"] + ["N m"] * 5 + ["%"] * 5 + ["mm", "-", "mm", "count", "count", "count"]
    count = lambda rs, flag: sum(r[col[flag]] for r in rs) if flag in col else 0
    srows = []
    for rs in groups:
        k = rs[0][col["step_no"]]
        tq = [max(abs(r[col[f"torque_A{i}"]]) for r in rs) for i in range(1, 6)]
        tp = [max(abs(r[col[f"torque_A{i}_of_limit"]]) for r in rs) for i in range(1, 6)]
        worst = min(rs, key=lambda r: r[col["self_clearance"]])
        pc = [r[col["part_clearance_from_arm"]] for r in rs if r[col["part_clearance_from_arm"]] != ""]
        srows.append([k, rs[0][col["step"]], rs[0][col["op"]], rs[0][col["motion"]], rs[0][col["t"]] - 1 / 240, rs[-1][col["t"]],
                      len(rs) / 240] + tq + tp + [worst[col["self_clearance"]], worst[col["self_closest_parts"]],
                                                  min(pc) if pc else "", count(rs, "gnn_called"), count(rs, "pinn_active"),
                                                  count(rs, "mgnn_called")])
    write(wb.create_sheet("Steps"), sh, su, srows, {"step": 24, "motion": 26, "closest_parts": 34})

    base = ["t", "step_no", "step"]
    for name, flag, prefix in (("GNN calls", "gnn_called", "gnn_"), ("PINN", "pinn_active", "pinn_"),
                               ("Motion GNN calls", "mgnn_called", "mgnn_")):
        if flag not in col or not any(r[col[flag]] == 1 for r in rows):
            continue
        keep = base + [h for h in head if h.startswith(prefix) and h != flag] + [f"q_A{i}" for i in range(1, 6)]
        idx = [col[h] for h in keep]
        write(wb.create_sheet(name), keep, [units[i] for i in idx], [[r[i] for i in idx] for r in rows if r[col[flag]] == 1],
              {"step": 24})

    cam = d["camera"]
    w, h, img = cam["w"], cam["h"], cam["img"]
    n = w * h
    look = next(r for r in rows if r[col["vision_called"]] == 1)
    ws = wb.create_sheet("Camera")
    ws.append(["The vision network's input: the camera image, %d x %d pixels, 3 colour channels (0-1). Its output: the part at "
               "x %.1f mm, z %.1f mm (really at x %.1f, z %.1f; error %.1f mm)." % (w, h, look[col["vision_out_x"]], look[col["vision_out_z"]],
                                                                                   look[col["vision_true_x"]], look[col["vision_true_z"]],
                                                                                   look[col["vision_error"]])])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["pixel_row", "pixel_col", "red", "green", "blue"])
    ws.append(["0 = top", "0 = left", "0-1", "0-1", "0-1"])
    for k in range(n):
        ws.append([k // w, k % w, img[k], img[n + k], img[2 * n + k]])
    for c in range(1, 6):
        ws.cell(3, c).font, ws.cell(3, c).fill = Font(bold=True), HEAD_FILL
        ws.cell(4, c).font, ws.cell(4, c).fill = Font(italic=True), UNIT_FILL
    ws.freeze_panes = "A5"

    out = HERE / args.out
    wb.save(out)
    print(f"wrote {out}: {len(rows)} control steps, {len(groups)} program lines")
    for r in srows:
        print(f"  {r[0]:2d} {r[1]:26s} {r[3]:28s} {r[4]:5.2f}-{r[5]:5.2f} s  peak {max(r[12:17]):5.1f} %  self {r[17]:6.1f} mm")


if __name__ == "__main__":
    main()
