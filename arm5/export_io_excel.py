"""Write the inputs and outputs of the 5-joint pose GNN and motion PINN to an Excel workbook.

Sheet "GNN": random targets (tool point, tool angle, payload), the 5 numbers the GNN reads, the 5 joint
angles it outputs, the angles after the page's two refinement steps, and what that pose means (reach
error, holding torque, clearances). Sheet "PINN": moves between GNN poses, sampled through time: the 13
numbers the PINN reads at each moment, the 5 numbers it outputs, and the joint angles, speeds,
accelerations and motor torques they give. Sheet "Columns" explains every column.

    python export_io_excel.py --gnn model_noself/arm5_gnn_joint.pt --pinn model_noself/arm5_pinn_joint.pt
"""
import argparse
import math
from pathlib import Path

import torch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import kin
from kin import (D, N, TAU_MAX, clearance, dls, fk, goal_features, holding, inverse_dynamics, sample_targets,
                 self_clearance)
from model import ChainGNN
from motion import T_REF, MotionPINN5, move_time

HERE = Path(__file__).parent
J = [f"A{i + 1}" for i in range(N)]


def load(gnn_file, pinn_file):
    ck = torch.load(HERE / gnn_file, map_location="cpu")
    gnn = ChainGNN(ck["hidden"], ck["steps"])
    gnn.load_state_dict(ck["state"])
    pinn = MotionPINN5()
    pinn.load_state_dict(torch.load(HERE / pinn_file, map_location="cpu")["state"])
    return gnn.eval(), pinn.eval()


def targets(n, g, random_poses=False):
    """Unseen targets that have a collision-free pose: the first 20,000 searched targets of dataset.pt, which the
    training scripts hold out. With random_poses, targets from random joint poses instead (some have none)."""
    if random_poses:
        return sample_targets(n, "cpu", g)[:3]
    d = torch.load(HERE / "dataset.pt")
    i = torch.randperm(20000, generator=g)[:n]
    return d["tcp"][i], d["phi"][i], d["payload"][i]


def gnn_rows(gnn, n, g, random_poses=False):
    tcp, phi, payload = targets(n, g, random_poses)
    f = goal_features(tcp, phi, payload)
    with torch.no_grad():
        raw = gnn(f)
    q = dls(raw, tcp, phi, steps=2).detach()
    err = lambda qq: (fk(qq)[0] - tcp).norm(dim=-1)
    hold = holding(q, payload) * 100
    cell, own = clearance(q, payload).min(-1).values, self_clearance(q).min(-1).values
    head = (["target", "x_mm", "z_mm", "tool_angle_deg", "payload_kg"]
            + ["in1_x", "in2_z", "in3_sin_angle", "in4_cos_angle", "in5_payload"]
            + [f"out_{j}_deg" for j in J] + [f"refined_{j}_deg" for j in J]
            + ["tool_error_network_mm", "tool_error_refined_mm"] + [f"holding_{j}_pct" for j in J]
            + ["clearance_cell_mm", "clearance_self_mm"])
    rows = []
    for i in range(n):
        rows.append([i + 1, tcp[i, 0].item(), tcp[i, 1].item(), phi[i].item() / D, payload[i].item()] + f[i].tolist()
                    + (raw[i] / D).tolist() + (q[i] / D).tolist() + [err(raw)[i].item(), err(q)[i].item()]
                    + hold[i].tolist() + [cell[i].item(), own[i].item()])
    return head, rows, (tcp, phi, payload, q)


def pinn_rows(pinn, poses, moves, k, g):
    tcp, phi, payload, q = poses
    i = torch.randint(len(q), (moves,), generator=g)
    j = torch.randint(len(q), (moves,), generator=g)
    head = (["move", "t_s", "s", "duration_T_s", "payload_kg"] + [f"q0_{a}_deg" for a in J] + [f"q1_{a}_deg" for a in J]
            + [f"in{c + 1}_q0_{a}" for c, a in enumerate(J)] + [f"in{c + 6}_dq_{a}" for c, a in enumerate(J)]
            + ["in11_time", "in12_log_T", "in13_payload"] + [f"out_N_{a}" for a in J]
            + [f"q_{a}_deg" for a in J] + [f"qd_{a}_deg_s" for a in J] + [f"qdd_{a}_deg_s2" for a in J]
            + [f"torque_{a}_pct" for a in J] + ["clearance_cell_mm", "clearance_self_mm"])
    rows = []
    s = torch.linspace(0, 1, k)
    for m in range(moves):
        q0, q1, p = q[i[m]], q[j[m]], payload[j[m]].item()      # start at one target's pose, carry target j's payload to its pose
        T = max(0.15, 1.5 * move_time(q0, q1).item())
        ex = lambda x: x[None].expand(k, *x.shape)
        Tk, pk = torch.full((k,), T), torch.full((k,), p)
        feats = torch.cat([ex((q0 - pinn.q_mid) / pinn.q_half), ex((q1 - q0) / pinn.q_half),
                           (2 * s - 1)[:, None], torch.full((k, 1), math.log(T / T_REF)), torch.full((k, 1), p / 12)], -1)
        with torch.no_grad():
            Nn = pinn.shape(ex(q0), ex(q1), s, Tk, pk)
        qq, qd, qdd = (x.detach() for x in pinn.motion(ex(q0), ex(q1), s, Tk, pk))
        tau = inverse_dynamics(qq, qd, qdd, pk) / TAU_MAX * 100
        cell, own = clearance(qq, pk).min(-1).values, self_clearance(qq).min(-1).values
        for t in range(k):
            rows.append([m + 1, s[t].item() * T, s[t].item(), T, p] + (q0 / D).tolist() + (q1 / D).tolist() + feats[t].tolist()
                        + Nn[t].tolist() + (qq[t] / D).tolist() + (qd[t] / D).tolist() + (qdd[t] / D).tolist()
                        + tau[t].tolist() + [cell[t].item(), own[t].item()])
    return head, rows


COLUMNS = [
    ("GNN", "target", "Row number."),
    ("GNN", "x_mm, z_mm", "Target tool point in the arm's plane (mm). The shoulder is at x 0, z 160."),
    ("GNN", "tool_angle_deg", "Target tool angle (deg, 0 = pointing along +x, -90 = straight down)."),
    ("GNN", "payload_kg", "Mass held at the tool point (kg)."),
    ("GNN", "in1..in5", "THE GNN'S INPUT: (x - 0)/780, (z - 160)/780, sin(angle), cos(angle), payload/12. "
                        "Placed on the goal node of a 7-node graph (5 joints, tool point, goal) whose edges carry link length/780."),
    ("GNN", "out_A1..A5_deg", "THE GNN'S OUTPUT: the 5 joint angles (deg), q = q_mid + q_half * tanh(y), so always within the limits "
                              "(A1 -20..200, A2-A5 -150..150)."),
    ("GNN", "refined_A1..A5_deg", "The same pose after the page's two damped-least-squares refinement steps."),
    ("GNN", "tool_error_network_mm / refined_mm", "Distance from the reached tool point to the target, before and after refinement."),
    ("GNN", "holding_A1..A5_pct", "Torque to hold the refined pose still with its payload, % of each motor's limit (kin.TAU_MAX)."),
    ("GNN", "clearance_cell_mm", "Closest gap from the arm to the floor, the bin and the carried part (negative = collision)."),
    ("GNN", "clearance_self_mm", "Closest gap between parts of the arm itself (links, housings, gripper) and its pedestal (negative = collision)."),
    ("PINN", "move", "Move number; each move is sampled at evenly spaced moments."),
    ("PINN", "t_s, s, duration_T_s", "Time into the move (s), normalized time s = t/T, and the move's duration T (1.5 x the joint blend's time)."),
    ("PINN", "q0_*, q1_*", "Start and end pose of the move (deg). Both are GNN poses; q1 is what the GNN chose for the target."),
    ("PINN", "in1..in5", "THE PINN'S INPUT, part 1: start pose normalized, (q0 - q_mid)/q_half."),
    ("PINN", "in6..in10", "THE PINN'S INPUT, part 2: pose change, (q1 - q0)/q_half."),
    ("PINN", "in11, in12, in13", "THE PINN'S INPUT, part 3: time 2s - 1, log(T / 1 s), payload/12."),
    ("PINN", "out_N_A1..A5", "THE PINN'S OUTPUT: the path correction N. Joint angles: q(s) = q0 + (q1 - q0) b(s) + w(s) N, "
                             "b(s) = 10s^3 - 15s^4 + 6s^5, w(s) = 64 s^3 (1-s)^3, so the move always starts at q0 and ends at q1."),
    ("PINN", "q_*, qd_*, qdd_*", "Joint angles (deg), speeds (deg/s) and accelerations (deg/s^2) at that moment, from the same call by autograd."),
    ("PINN", "torque_A1..A5_pct", "Motor torque needed at that moment (Newton-Euler), % of each motor's limit."),
    ("PINN", "clearance_cell_mm, clearance_self_mm", "As on the GNN sheet, at that moment of the move."),
]


def unit(name):
    """Unit of a column, from its name."""
    if name.startswith("in"):
        return "unitless (normalized)"
    for suffix, u in (("_deg_s2", "deg/s^2"), ("_deg_s", "deg/s"), ("_deg", "deg"), ("_mm", "mm"), ("_kg", "kg"),
                      ("_pct", "% of motor limit"), ("_s", "s")):
        if name.endswith(suffix):
            return u
    return {"s": "t/T, unitless"}.get(name, "rad" if name.startswith("out_N") else "-")


def sheet(ws, head, rows, decimals=3):
    ws.append(head)
    ws.append([unit(h) for h in head])
    for r in rows:
        ws.append([round(v, decimals) if isinstance(v, float) else v for v in r])
    fill = PatternFill("solid", fgColor="DDE6F0")
    for c in range(1, len(head) + 1):
        cell = ws.cell(1, c)
        cell.font, cell.fill, cell.alignment = Font(bold=True), fill, Alignment(wrap_text=True, vertical="top")
        u = ws.cell(2, c)
        u.font, u.fill, u.alignment = Font(italic=True), PatternFill("solid", fgColor="F1F4F8"), Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(16, len(head[c - 1]) + 2))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 45, 30
    ws.freeze_panes = "B3"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--gnn", default="arm5_gnn.pt")
    ap.add_argument("--pinn", default="arm5_pinn.pt")
    ap.add_argument("--targets", type=int, default=1000)
    ap.add_argument("--moves", type=int, default=50)
    ap.add_argument("--moments", type=int, default=21)
    ap.add_argument("--out", default="model_io.xlsx")
    ap.add_argument("--random-poses", action="store_true", help="targets from random joint poses, reachable or not")
    args = ap.parse_args()
    torch.set_grad_enabled(True)
    g = torch.Generator().manual_seed(2026)
    gnn, pinn = load(args.gnn, args.pinn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Columns"
    ws.append([f"Inputs and outputs of the 5-joint arm's networks: GNN {args.gnn}, PINN {args.pinn}. Motors "
               f"{', '.join(f'{v:g}' for v in kin.TAU_MAX.tolist())} N m. Targets: "
               + ("from random joint poses (some have no collision-free pose)" if args.random_poses else
                  "unseen searched targets, each reachable without a collision")])
    ws.append([])
    ws.append(["Row 2 of each data sheet gives every column's unit. The network inputs (in...) are normalized numbers without units."])
    ws.append(["Sheet", "Column", "Meaning"])
    for r in COLUMNS:
        ws.append(list(r))
    ws["A1"].font = Font(bold=True, size=12)
    for c in "ABC":
        ws[f"{c}4"].font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 8, 34, 120
    head, rows, poses = gnn_rows(gnn, args.targets, g, args.random_poses)
    sheet(wb.create_sheet("GNN"), head, rows)
    head, rows = pinn_rows(pinn, poses, args.moves, args.moments, g)
    sheet(wb.create_sheet("PINN"), head, rows)
    wb.save(HERE / args.out)
    print(f"wrote {HERE / args.out}: {args.targets} GNN targets, {args.moves} PINN moves x {args.moments} moments")


if __name__ == "__main__":
    main()
